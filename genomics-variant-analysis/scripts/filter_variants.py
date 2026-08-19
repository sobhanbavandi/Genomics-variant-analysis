#!/usr/bin/env python3
"""
filter_variants.py
===================

Reproducible variant-prioritization pipeline for a single-sample ANNOVAR
"multianno" annotation table (table_annovar.pl output, hg19), applied to a
consanguineous-family case with a severe neurodevelopmental / cardiac
phenotype.

WHY THIS FILTER LOGIC
----------------------
A raw exome/genome VCF annotated with ANNOVAR against dozens of databases
routinely contains hundreds of thousands of variants. The overwhelming
majority are common, functionally silent, or irrelevant to a rare Mendelian
phenotype. Clinical/research exome analysis narrows this down with a
standard, defensible funnel:

  1. Quality  -> keep only variants that passed the caller's own filter
                 (VCF FILTER == PASS)
  2. Location -> keep only variants that can plausibly change a protein
                 (exonic or ANNOVAR-annotated splicing regions of RefSeq genes;
                 excludes intronic/intergenic/UTR/ncRNA)
  3. Effect   -> drop synonymous changes, which do not alter the protein
  4. Rarity   -> keep only variants that are rare or entirely absent from
                 five independent population-frequency panels. The default
                 threshold is <=0.1%; a broader <=1% exploratory pass can be
                 requested from the command line.
  5. Zygosity -> keep only HOMOZYGOUS-ALT genotypes in the proband

The zygosity step is the case-specific piece: the family history describes multiple
generations of consanguinity and a single affected child with no family
history of the same condition. That pattern is the textbook signature of
autosomal recessive disease transmitted through a shared ancestral haplotype
("autozygosity") -- so, in the absence of parental sequencing to build a
trio, homozygosity in the proband is the most powerful single filter
available. (Its main blind spot -- compound heterozygosity -- is noted in
the README under Limitations.)

USAGE
-----
    python3 filter_variants.py /path/to/annovar_multianno.xlsx results/candidates_homozygous_rare_coding.csv
    # broader exploratory pass:
    python3 filter_variants.py /path/to/annovar_multianno.xlsx results/candidates_homozygous_rare_coding_broad_1pct.csv --maf-threshold 0.01

Requires: openpyxl, pandas (both pip-installable; see requirements.txt)

INPUT ASSUMPTIONS
------------------
The script expects the standard column layout produced by ANNOVAR's
table_annovar.pl when run with `-otherinfo` against a VCF converted with
convert2annovar.pl -includeinfo. If your file was annotated with a
different -protocol list, the column *positions* below will not line up --
open the file once and adjust the `COL` dictionary to match your header.
"""

import argparse
import csv
import sys
import time
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

# --------------------------------------------------------------------------
# Column positions (0-indexed) in this project's multianno file.
# Re-check these against your own header row if you re-use this script on a
# different ANNOVAR run -- protocol order changes the column layout.
# --------------------------------------------------------------------------
COL = dict(
    chr=0, start=1, end=2, ref=3, alt=4,
    func=5, gene=6, genedetail=7, exonicfunc=8, aachange=9,
    sift_pred=13, pp2_hdiv_pred=15, pp2_hvar_pred=17,
    cadd_phred=36,
    esp6500=58, exac_all=59,
    kaviar_af=67,
    hrc_af=70,
    thousandg_all=76,
    avsnp=81,
    clinsig=82, clndbn=83,
    intervar=87,
    vcf_filter=125, vcf_format=127, vcf_sample=128,
)

STANDARD_CHRS = {f"chr{i}" for i in range(1, 23)} | {"chrX", "chrY", "chrM"}
LOF_TERMS = {"stopgain", "stoploss", "frameshift deletion", "frameshift insertion"}
RARE_MAF_THRESHOLD = 0.001  # 0.1%; conservative default for severe rare Mendelian disease


def is_rare_or_absent(val, thresh=RARE_MAF_THRESHOLD):
    """A '.' in ANNOVAR output means 'not observed in this database' --
    treated here as consistent with a rare variant, per standard practice."""
    if val in (None, ".", ""):
        return True
    try:
        return float(val) <= thresh
    except (ValueError, TypeError):
        return True


def get_gt(fmt, sample):
    """Pull the GT sub-field out of a VCF FORMAT:SAMPLE pair, e.g.
    fmt='GT:AD:DP:GQ:PL', sample='1/1:0,27:27:81:973,81,0' -> '1/1'."""
    if not fmt or not sample:
        return None
    fields, vals = fmt.split(":"), sample.split(":")
    if "GT" not in fields or len(vals) < len(fields):
        return None
    return vals[fields.index("GT")].replace("|", "/")


def zygosity(gt):
    if gt is None:
        return "unknown"
    alleles = gt.split("/")
    if len(alleles) != 2:
        return "other"
    a, b = alleles
    if a == "." or b == ".":
        return "missing"
    if a == b:
        return "hom_ref" if a == "0" else "hom_alt"
    return "het"


def classify_tier(exfunc, is_splicing, intervar, clinsig, sift, pp2h, pp2v, cadd_raw):
    """A simple, transparent priority score -- NOT a clinical classification.
    Tier 1: loss-of-function, or already flagged pathogenic/likely
            pathogenic by InterVar/ClinVar in this (older) annotation.
    Tier 2: missense with >=2 of 4 in-silico tools calling it damaging.
    Tier 3: everything else that survived the upstream filters.
    """
    intervar = (intervar or ".").strip().lower()
    clinsig = (clinsig or ".").strip().lower()
    try:
        cadd = float(cadd_raw)
    except (TypeError, ValueError):
        cadd = None

    if "pathogenic" in intervar or "pathogenic" in clinsig:
        return 1
    if exfunc in LOF_TERMS or (is_splicing and exfunc in (".", "unknown")):
        return 1
    if exfunc == "nonsynonymous SNV":
        votes = sum([
            sift == "D",
            pp2h in ("D", "P"),
            pp2v in ("D", "P"),
            cadd is not None and cadd >= 20,
        ])
        return 2 if votes >= 2 else 3
    return 3


def run(in_path: str, out_csv: str, maf_threshold: float = RARE_MAF_THRESHOLD):
    t0 = time.time()
    wb = load_workbook(in_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    next(rows)  # header

    counts = Counter()
    candidates = []

    for row in rows:
        counts["total"] += 1

        if row[COL["chr"]] not in STANDARD_CHRS:
            continue
        counts["standard_chr"] += 1

        if row[COL["vcf_filter"]] != "PASS":
            continue
        counts["pass_filter"] += 1

        func = row[COL["func"]] or ""
        tokens = set(func.split(";"))
        is_exonic = "exonic" in tokens        # exact match excludes ncRNA_exonic
        is_splicing = "splicing" in tokens    # exact match excludes ncRNA_splicing
        if not (is_exonic or is_splicing):
            continue
        counts["coding_or_splicing"] += 1

        exfunc = row[COL["exonicfunc"]] or "."
        if exfunc == "synonymous SNV":
            continue
        if exfunc == "unknown" and not is_splicing:
            continue
        counts["protein_altering"] += 1

        rare = all(is_rare_or_absent(row[COL[c]], maf_threshold) for c in
                   ("thousandg_all", "exac_all", "esp6500", "kaviar_af", "hrc_af"))
        if not rare:
            continue
        counts["rare"] += 1

        gt = get_gt(row[COL["vcf_format"]], row[COL["vcf_sample"]])
        zyg = zygosity(gt)
        if zyg != "hom_alt":
            continue
        counts["hom_alt"] += 1

        tier = classify_tier(
            exfunc, is_splicing, row[COL["intervar"]], row[COL["clinsig"]],
            row[COL["sift_pred"]], row[COL["pp2_hdiv_pred"]], row[COL["pp2_hvar_pred"]],
            row[COL["cadd_phred"]],
        )

        candidates.append({
            "Tier": tier,
            "Chr": row[COL["chr"]], "Start": row[COL["start"]], "End": row[COL["end"]],
            "Ref": row[COL["ref"]], "Alt": row[COL["alt"]],
            "Func.refGene": func, "Gene.refGene": row[COL["gene"]],
            "GeneDetail.refGene": row[COL["genedetail"]],
            "ExonicFunc.refGene": exfunc, "AAChange.refGene": row[COL["aachange"]],
            "SIFT_pred": row[COL["sift_pred"]],
            "Polyphen2_HDIV_pred": row[COL["pp2_hdiv_pred"]],
            "Polyphen2_HVAR_pred": row[COL["pp2_hvar_pred"]],
            "CADD_phred": row[COL["cadd_phred"]],
            "1000g2015aug_all": row[COL["thousandg_all"]], "ExAC_ALL": row[COL["exac_all"]],
            "esp6500siv2_all": row[COL["esp6500"]], "Kaviar_AF": row[COL["kaviar_af"]],
            "HRC_AF": row[COL["hrc_af"]], "avsnp147": row[COL["avsnp"]],
            "CLINSIG": row[COL["clinsig"]], "CLNDBN": row[COL["clndbn"]],
            "InterVar_automated": row[COL["intervar"]],
            "Genotype": gt, "Zygosity": zyg,
        })

    wb.close()
    candidates.sort(key=lambda r: (r["Tier"], r["Chr"], r["Start"]))

    Path(out_csv).parent.mkdir(parents=True, exist_ok=True)
    if candidates:
        with open(out_csv, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(candidates[0].keys()))
            w.writeheader()
            w.writerows(candidates)

    print("==== FILTERING FUNNEL ====")
    for label, key in [
        ("Total variants (SNPs+indels)", "total"),
        ("Standard chromosomes (1-22,X,Y,M)", "standard_chr"),
        ("PASS quality filter", "pass_filter"),
        ("Exonic / ANNOVAR-annotated splicing", "coding_or_splicing"),
        ("Protein-altering (non-synonymous)", "protein_altering"),
        (f"Rare/absent in 5 population DBs (<={maf_threshold:g})", "rare"),
        ("Homozygous-ALT in proband", "hom_alt"),
    ]:
        print(f"  {label:45s} {counts[key]:>8,}")
    print(f"\nWrote {len(candidates)} candidate variants -> {out_csv}")
    print(f"Elapsed: {time.time() - t0:.1f}s")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Filter an ANNOVAR multianno XLSX for rare homozygous coding/splicing candidates.")
    parser.add_argument("in_file", help="Path to the ANNOVAR multianno XLSX")
    parser.add_argument("output", nargs="?", default="results/candidates_homozygous_rare_coding.csv", help="Output CSV path")
    parser.add_argument("--maf-threshold", type=float, default=RARE_MAF_THRESHOLD, help="Maximum population AF accepted in every checked database (default: 0.001 = 0.1%%; use 0.01 for the broader exploratory pass)")
    args = parser.parse_args()
    if not (0 <= args.maf_threshold <= 1):
        parser.error("--maf-threshold must be between 0 and 1")
    run(args.in_file, args.output, args.maf_threshold)
